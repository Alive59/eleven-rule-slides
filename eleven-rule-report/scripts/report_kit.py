"""Report-page helpers for eleven-rule-report (python-pptx).

Native charts in the red-black style and the comparison table style, to be used on
slides made by ../../scripts/slide_kit.py. Set PALETTE from the deck's term sheet.
"""
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
import copy
from lxml import etree

# Red-black style. Colors come from the term sheet, bound to display names.
PALETTE = {"前回": "B2241E", "今回": "222222", "参照": "7F7F7F"}
GRID, AXIS_LINE, AXIS_TEXT = "E6E6E6", "BFBFBF", "595959"
AXIS_PT, LABEL_PT = 14, 16

def _ea(font, face):
    font._rPr.append(etree.Element(qn("a:ea"), typeface=face))

def _style_common(chart, number_format, show_value_axis=True, font_en="Calibri", font_ja="Yu Gothic"):
    chart.has_title = False
    chart.has_legend = False
    chart.font.size = Pt(AXIS_PT); chart.font.name = font_en; _ea(chart.font, font_ja)
    ca, va = chart.category_axis, chart.value_axis
    ca.has_title = False; va.has_title = False
    ca.format.line.color.rgb = RGBColor.from_string(AXIS_LINE)
    ca.tick_labels.font.bold = True; ca.tick_labels.font.size = Pt(AXIS_PT)
    ca.tick_labels.font.color.rgb = RGBColor(0, 0, 0)
    if show_value_axis:
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRID)
        va.format.line.color.rgb = RGBColor.from_string(AXIS_LINE)
        va.tick_labels.font.size = Pt(AXIS_PT)
        va.tick_labels.font.color.rgb = RGBColor.from_string(AXIS_TEXT)
        va.tick_labels.number_format = "General"; va.tick_labels.number_format_is_linked = False
    else:
        va.visible = False; va.has_major_gridlines = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(LABEL_PT); dl.font.bold = True; dl.font.color.rgb = RGBColor(0, 0, 0)
    # Write the format out. number_format_is_linked=True alone makes python-pptx emit
    # <c:numFmt sourceLinked="1"/> without the required formatCode -> PowerPoint repair.
    dl.number_format = number_format; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    _positive_axis_ids(chart)
    return plot


def _positive_axis_ids(chart):
    """python-pptx's chart template uses negative axis IDs; the schema says unsignedInt.
    Remap them to positive values everywhere they are referenced."""
    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    root = chart._chartSpace
    ids = {}
    for el in root.iter(C + "axId", C + "crossAx"):
        v = el.get("val")
        if v is not None and v.startswith("-"):
            ids.setdefault(v, str(500000000 + len(ids) + 1))
            el.set("val", ids[v])

def compare_bars(slide, x, y, w, h, labels, values, number_format="0.0", y_min=None):
    """One metric, one column per term (前回/今回/参照...), left to right. Colors by label."""
    cd = CategoryChartData(number_format=number_format)
    cd.categories = labels
    cd.add_series("v", values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w), Inches(h), cd).chart
    plot = _style_common(chart, number_format); plot.gap_width = 50
    if y_min is not None:
        chart.value_axis.minimum_scale = y_min
    for i, lab in enumerate(labels):
        f = plot.series[0].points[i].format.fill
        f.solid(); f.fore_color.rgb = RGBColor.from_string(PALETTE[lab])
    return chart

def grouped_bars(slide, x, y, w, h, categories, series, number_format="0.0"):
    """Terms as series per category (e.g. bins), value axis hidden, labels on bars.
    series = {"前回": [...], "今回": [...]}. Pair with slide_legend() once per slide."""
    cd = CategoryChartData(number_format=number_format)
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, vals)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w), Inches(h), cd).chart
    plot = _style_common(chart, number_format, show_value_axis=False); plot.gap_width = 60; plot.overlap = 0
    plot.data_labels.font.size = Pt(AXIS_PT); plot.data_labels.font.bold = False
    for s, name in zip(plot.series, series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = RGBColor.from_string(PALETTE[name])
    return chart

def slide_legend(slide, x, y, names, font_en="Calibri", font_ja="Yu Gothic", size=20):
    """One shared legend row on the slide: swatch + display name per term."""
    from pptx.enum.shapes import MSO_SHAPE
    cx = x
    for name in names:
        sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(y + 0.1), Inches(0.6), Inches(0.22))
        plain_fill(sw, PALETTE[name])
        tb = slide.shapes.add_textbox(Inches(cx + 0.7), Inches(y - 0.05), Inches(1.6), Inches(0.5))
        r = tb.text_frame.paragraphs[0].add_run(); r.text = name
        r.font.size = Pt(size); r.font.name = font_en; _ea(r.font, font_ja)
        cx += 2.4

def label_under(slide, x, w, text, y=6.65, font_en="Calibri", font_ja="Yu Gothic", size=20):
    """20pt bold label centered under a chart or panel (replaces in-chart titles)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.45))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True
    r.font.name = font_en; _ea(r.font, font_ja)

def _set_borders(cell, spec):
    """spec = {"L": (pt, "RRGGBB") or None, ...}; written in schema order, before the fill."""
    tcPr = cell._tc.get_or_add_tcPr()
    for side in "LRTB":
        for old in tcPr.findall(qn(f"a:ln{side}")):
            tcPr.remove(old)
    for i, side in enumerate("LRTB"):
        w_pt, color = spec.get(side) or (0, None)
        ln = etree.Element(qn(f"a:ln{side}"), w=str(int(w_pt * 12700)))
        if w_pt:
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"), val=color)
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)

def report_table(slide, x, y, w, h, row_label, groups, metrics, rows, bold=(),
                 font_en="Calibri", font_ja="Yu Gothic", size=14):
    """Page-9 style comparison table: white, no fills, bold 2-row header (group over metric),
    black 1pt vertical rule after the label column and between groups, 1pt #595959 rule under
    the header and the last (全体/ALL) row, 0.75pt #D9D9D9 rules between body rows.
    rows = [("≤16", [g1m1, g1m2, ..., g2m1, ...]), ..., ("全体", [...])]
    bold = {(body_row, value_col)} — the better value of each pair."""
    nm = len(metrics); ncol = 1 + nm * len(groups); nrow = 2 + len(rows)
    gf = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table; tbl.first_row = False; tbl.horz_banding = False
    def put(r, c, text, b=False, align=PP_ALIGN.CENTER):
        cell = tbl.cell(r, c); cell.fill.background(); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; p.alignment = align
        if text != "":
            run = p.add_run(); run.text = str(text)
            run.font.size = Pt(size); run.font.bold = b; run.font.name = font_en
            run.font._rPr.append(etree.Element(qn("a:ea"), typeface=font_ja))
    put(0, 0, ""); put(1, 0, row_label, True, PP_ALIGN.LEFT)
    for g, name in enumerate(groups):
        c0 = 1 + g * nm
        tbl.cell(0, c0).merge(tbl.cell(0, c0 + nm - 1)); put(0, c0, name, True)
        for m, met in enumerate(metrics):
            put(1, c0 + m, met, True)
    for r, (lab, vals) in enumerate(rows):
        put(2 + r, 0, lab, True, PP_ALIGN.LEFT)
        for c, v in enumerate(vals):
            put(2 + r, 1 + c, v, (r, c) in set(bold))
    for r in range(nrow):
        for c in range(ncol):
            spec = {}
            if c == 0 or (c % nm == 0 and c != ncol - 1):
                spec["R"] = (1.0, "000000")
            if c > 1 and (c - 1) % nm == 0:          # merged group headers need the left rule too
                spec["L"] = (1.0, "000000")
            if r == 1 or r == nrow - 1:
                spec["B"] = (1.0, "595959")
            elif r >= 2:
                spec["B"] = (0.75, "D9D9D9")
            _set_borders(tbl.cell(r, c), spec)
    return tbl

# --- Flat shapes: color fill only -------------------------------------------------
# Autoshapes carry a <p:style> that links them to the theme: fillRef idx=3 is a
# gradient and effectRef idx=2 a shadow in the default theme. An empty effectLst is
# not enough (some renderers still draw the theme effect), so the link is removed and
# its colors are written out explicitly.

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_FILLS = {_A + t for t in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")}
_COLORS = {_A + t for t in ("schemeClr", "srgbClr", "prstClr", "sysClr", "scrgbClr", "hslClr")}


def _color_of(ref):
    """Copy of the color element inside a style ref (fillRef/lnRef/fontRef), or None."""
    if ref is None:
        return None
    for c in ref:
        if c.tag in _COLORS:
            return copy.deepcopy(c)
    return None


def _solid(color_el):
    f = etree.Element(_A + "solidFill")
    f.append(color_el)
    return f


def _flatten(el):
    """el is a p:sp or p:cxnSp. Leaves: solid fill (filled shapes, no outline) or a
    plain line (unfilled shapes and connectors). No style link, effects, 3D, gradient."""
    spPr = el.find(_P + "spPr")
    style = el.find(_P + "style")
    is_cxn = etree.QName(el).localname == "cxnSp"
    fill_ref = _color_of(style.find(_A + "fillRef")) if style is not None else None
    ln_ref = _color_of(style.find(_A + "lnRef")) if style is not None else None
    font_ref = _color_of(style.find(_A + "fontRef")) if style is not None else None

    for tag in ("effectLst", "effectDag", "scene3d", "sp3d"):
        for x in spPr.findall(_A + tag):
            spPr.remove(x)
    fills = [c for c in spPr if c.tag in _FILLS]
    fill = fills[0] if fills else None
    if fill is not None and fill.tag in (_A + "gradFill", _A + "pattFill"):
        stop = fill.find(".//" + _A + "gs")       # first gradient stop -> solid
        col = next((copy.deepcopy(c) for c in (stop if stop is not None else fill)
                    if c.tag in _COLORS), None)
        new = _solid(col) if col is not None else etree.Element(_A + "noFill")
        fill.addprevious(new); spPr.remove(fill); fill = new
    if fill is None:
        if is_cxn or style is None or fill_ref is None:
            fill = etree.Element(_A + "noFill")
        else:
            fill = _solid(fill_ref)
        geom = [c for c in spPr if c.tag in (_A + "xfrm", _A + "prstGeom", _A + "custGeom")]
        (geom[-1].addnext(fill) if geom else spPr.insert(0, fill))
    filled = fill.tag != _A + "noFill"

    ln = spPr.find(_A + "ln")
    if ln is None:
        ln = etree.Element(_A + "ln")
    for x in ln.findall(_A + "effectLst") + ln.findall(_A + "effectDag"):
        ln.remove(x)
    if filled and not is_cxn:                     # filled shape: no outline at all
        for child in list(ln):
            ln.remove(child)
        etree.SubElement(ln, _A + "noFill")
    elif not any(c.tag in _FILLS for c in ln):    # line shape relying on the theme color
        ln.insert(0, _solid(ln_ref if ln_ref is not None else etree.Element(_A + "srgbClr", val="000000")))
    fill.addnext(ln)
    ln.addnext(etree.Element(_A + "effectLst"))

    if style is not None:
        if font_ref is not None:                  # keep the text color the style gave
            for rpr in el.iter(_A + "rPr", _A + "endParaRPr"):
                if not any(c.tag in _FILLS for c in rpr):
                    rpr.insert(0, _solid(copy.deepcopy(font_ref)))
            for r in el.iter(_A + "r"):
                if r.find(_A + "rPr") is None:
                    rpr = etree.Element(_A + "rPr"); r.insert(0, rpr)
                    rpr.append(_solid(copy.deepcopy(font_ref)))
        el.remove(style)


def plain_fill(shape, color):
    """Solid color fill and nothing else: no outline, shadow, glow, bevel or gradient."""
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
    _flatten(shape._element)
    return shape


def _shape_elements(prs):
    for i, slide in enumerate(prs.slides, 1):
        for el in slide.shapes._spTree.iter(_P + "sp", _P + "cxnSp"):
            if el.find(".//" + _P + "cNvSpPr[@txBox='1']") is not None:
                continue                          # text boxes carry no shape styling
            if el.find(_P + "spPr") is None:
                continue
            yield i, el


def flatten_shapes(prs):
    """Pass over a whole deck before saving: every autoshape and connector ends with a
    color fill only (filled shapes) or a plain line only (red enclosures, dashed frames,
    arrows). Pictures, charts, tables and text boxes are left alone. Returns the count."""
    n = 0
    for _, el in _shape_elements(prs):
        _flatten(el); n += 1
    return n


def flat_violations(prs):
    """[(slide_no, shape_name, reason)] for shapes styled beyond a color fill."""
    out = []
    for i, el in _shape_elements(prs):
        name = el.find(".//" + _P + "cNvPr").get("name")
        spPr = el.find(_P + "spPr")
        if el.find(_P + "style") is not None:
            out.append((i, name, "linked to theme styles (shadow/gradient may apply)"))
        for tag in ("scene3d", "sp3d", "gradFill", "pattFill"):
            if spPr.find(_A + tag) is not None:
                out.append((i, name, f"has {tag}"))
        eff = spPr.find(_A + "effectLst")
        if (eff is not None and len(eff)) or spPr.find(_A + "effectDag") is not None:
            out.append((i, name, "has shadow/glow/soft edge"))
        fills = [c for c in spPr if c.tag in _FILLS]
        ln = spPr.find(_A + "ln")
        if fills and fills[0].tag == _A + "solidFill" and etree.QName(el).localname == "sp" \
                and ln is not None and ln.find(_A + "noFill") is None:
            out.append((i, name, "filled shape has an outline"))
    return out


# --- Chart check ---------------------------------------------------------------------
# Defects PowerPoint repairs on open that a slide-level check misses, because charts
# live in their own parts.

_SER_ORDER = ["idx", "order", "tx", "spPr", "invertIfNegative", "pictureOptions", "dPt",
              "dLbls", "trendline", "errBars", "cat", "val", "shape", "extLst"]


def chart_violations(prs):
    """[(slide_no, chart_part, reason)] for chart XML PowerPoint will repair."""
    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not getattr(sh, "has_chart", False):
                continue
            root = sh.chart._chartSpace
            part = sh.chart.part.partname
            for nf in root.iter(C + "numFmt"):
                if nf.get("formatCode") is None:
                    out.append((i, part, "numFmt without formatCode"))
            for el in root.iter(C + "axId", C + "crossAx"):
                if el.get("val", "0").startswith("-"):
                    out.append((i, part, f"negative axis id {el.get('val')}"))
            for ser in root.iter(C + "ser"):
                pos = [_SER_ORDER.index(etree.QName(c).localname) for c in ser
                       if etree.QName(c).localname in _SER_ORDER]
                if pos != sorted(pos):
                    out.append((i, part, "series children out of order (e.g. dPt after dLbls)"))
    return out
