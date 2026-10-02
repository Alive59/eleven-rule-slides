#!/usr/bin/env python3
"""Audit a .pptx against the mechanical parts of the eleven rules.

    python check_rules.py deck.pptx [--json]

FAIL = a rule is broken and the deck should not ship.
WARN = likely a problem; needs a look or a stated reason.

Zones are inferred per slide, not assumed: the topmost text shape is the title, wide
text below it and above the first figure is the summary block, and everything else
is annotation. Diagrams, flow charts, grouped shapes and figure annotations are
exempt from the font and size rules. That way the audit works on decks not built with
slide_kit.js.

Content rules (1 wording, 4 keywords, 6 conclusion, 7 numbers, 9 separation) still
need a human read — this only covers what geometry and XML can prove.
"""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_IN = 914400.0
# Overridden by --font-en / --font-ja; the deck's fonts are the user's choice, the
# rule is only that the deck is consistent about them.
LATIN_FONTS = {"Calibri"}
JA_FONTS = {"Yu Gothic", "Yu Gothic UI", "游ゴシック", "游ゴシック Medium", "Yu Gothic Medium"}
JA_VARIANTS = (" UI", " Medium", " Light", " Bold")
CJK_RE = re.compile(r"[\u3040-\u309F\u30A0-\u30FF\u3400-\u4DBF\u4E00-\u9FFF]")
LATIN_RE = re.compile(r"[A-Za-z]")
RESULT_TITLE_RE = re.compile(
    r"result|accuracy|evaluat|experiment|dataset|data set|performance|problem|"
    r"comparison|ablation|結果|評価|実験|精度|データ|性能", re.I)
GENERIC_TITLE_RE = re.compile(
    r"^[\d.\s]*(results?|discussion|method|conclusion|introduction|experiments?|"
    r"background|overview|summary|まとめ|結果|考察|手法|概要|背景)\s*(\[\d/\d\])?$", re.I)
# Rule 4: a title is a compact keyword phrase, never a sentence. Keep in sync with
# slide_kit.py / slide_kit.js.
TITLE_TERMINAL_RE = re.compile(r"[.。!！?？]\s*$")
TITLE_JA_PREDICATE_RE = re.compile(
    r"(です|ます|ました|ません|でした|である|だった|した|する|される|された|"
    r"できる|できた|いる|ない|なった|なる)[。.!！?？]?\s*$")
TITLE_EN_VERB_RE = re.compile(
    r"\b(is|are|was|were|has|have|had|does|do|did|can|could|will|would|should|"
    r"we|our|improves?|improved|outperforms?|outperformed|reduces?|reduced|"
    r"increases?|increased|decreases?|decreased|achieves?|achieved|shows?|showed|"
    r"enables?|enabled|yields?|yielded|reaches|reached|fails|failed|leads|led)\b",
    re.I)
MAX_TITLE_WORDS = 8     # Latin words in a compact title
MAX_TITLE_JA_CHARS = 20 # characters in a compact Japanese title

TITLE_SEARCH_Y = 1.60   # in — the title must start within this of the top
FALLBACK_EVIDENCE_Y = 2.75
WIDE_FRAC = 0.40        # a summary box spans at least this fraction of the slide
PAGE_NUM_PT = 28
MIN_TITLE_PT = 32
MIN_BODY_PT = 20
MIN_ANY_PT = 14
EDGE_MARGIN = 0.40
LATIN_TAG = re.compile(r'<a:latin[^>]*typeface="([^"]*)"')
EA_TAG = re.compile(r'<a:ea[^>]*typeface="([^"]*)"')


def inches(v):
    return None if v is None else v / EMU_IN


def iter_shapes(shapes, in_group=False):
    """Yield (shape, in_group). Anything inside a group is part of a diagram or
    flow chart, so it is exempt from the font and size rules."""
    for sh in shapes:
        yield sh, in_group
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(sh, "shapes"):
            for sub, _ in iter_shapes(sh.shapes, True):
                yield sub, True


def is_visual(sh):
    return (
        sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE,
                          MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.CHART)
        or getattr(sh, "has_chart", False)
        or getattr(sh, "has_table", False)
    )


def run_fonts(run):
    """(latin_typeface, ea_typeface) as written on the run, or (None, None).

    python-pptx's font.name only exposes <a:latin>; Japanese glyphs are rendered
    with <a:ea>, so a correctly styled mixed run looks wrong if you read only one.
    """
    xml = run._r.xml
    lat = LATIN_TAG.search(xml)
    ea = EA_TAG.search(xml)
    return (lat.group(1) if lat else None, ea.group(1) if ea else None)


SZ_TAG = re.compile(r'sz="(\d+)"')
CLR_TAG = re.compile(r'<a:srgbClr val="([0-9A-Fa-f]{6})"')


def page_number_style(slide, pn):
    """(pt, rrggbb) for the slide-number field, falling back to the layout and
    master, since a deck may style the number once on the master."""
    sources = [pn._element.xml]
    for part in (slide.slide_layout, slide.slide_layout.slide_master):
        for sh in part.shapes:
            if "slidenum" in sh._element.xml:
                sources.append(sh._element.xml)
    size = color = None
    for xml in sources:
        fld = re.search(r"<a:fld.*?</a:fld>", xml, re.S)
        # run properties on the field win; the placeholder's list style is the
        # fallback, which is where a master-defined page number keeps its size
        for scope in ([fld.group(0)] if fld else []) + [xml]:
            if size is None:
                m = SZ_TAG.search(scope)
                if m:
                    size = int(m.group(1)) / 100.0
            if color is None:
                m = CLR_TAG.search(scope)
                if m:
                    color = m.group(1)
    return size, color


def has_page_number(slide):
    for sh in slide.shapes:
        if "slidenum" in sh._element.xml:
            return sh
    return None


def check_slide(slide, slide_w, slide_h):
    issues = []

    def fail(rule, msg):
        issues.append(("FAIL", rule, msg))

    def warn(rule, msg):
        issues.append(("WARN", rule, msg))

    shapes = list(iter_shapes(slide.shapes))
    grouped = {id(sh._element) for sh, g in shapes if g}
    shapes = [sh for sh, _ in shapes]
    # --- zone inference -----------------------------------------------------
    pn = has_page_number(slide)
    pn_el = pn._element if pn is not None else None
    texts = [sh for sh in shapes
             if sh.has_text_frame and sh.text_frame.text.strip()
             and sh._element is not pn_el and sh.top is not None]
    visuals = [sh for sh in shapes if is_visual(sh) and sh.top is not None]
    evidence_top = min((inches(v.top) for v in visuals), default=None)
    if evidence_top is None or evidence_top < 0.8:
        evidence_top = FALLBACK_EVIDENCE_Y

    title_sh = None
    if texts:
        top_sh = min(texts, key=lambda s: s.top)
        if inches(top_sh.top) < TITLE_SEARCH_Y:
            title_sh = top_sh
    title_bottom = (inches(title_sh.top) + inches(title_sh.height)) if title_sh else 1.10

    def zone_of(sh):
        if sh is title_sh:
            return "title"
        top = inches(sh.top)
        wide = inches(sh.width) >= WIDE_FRAC * slide_w
        if top < max(evidence_top, title_bottom) - 0.05 and wide:
            return "summary"
        return "annotation"

    # --- per-shape checks ---------------------------------------------------
    summary_paras, big_text = [], []
    numeric_in_summary = False

    for sh in shapes:
        left, top = inches(sh.left), inches(sh.top)
        width, height = inches(sh.width), inches(sh.height)
        if None not in (left, top, width, height):
            if left < -0.01 or top < -0.01 or left + width > slide_w + 0.01 or top + height > slide_h + 0.01:
                fail("2 placement", f"{sh.shape_type} extends outside the slide")
            elif sh._element is not pn_el and (left < EDGE_MARGIN or top < 0.2
                                   or left + width > slide_w - EDGE_MARGIN
                                   or top + height > slide_h - 0.25):
                warn("2 placement", f"shape closer than {EDGE_MARGIN}\" to a slide edge")

        if not sh.has_text_frame or sh._element is pn_el:
            continue
        zone = zone_of(sh)
        # Diagrams, flow charts and figure annotations are excluded from the font
        # and size rules — their text is part of the artwork, not slide copy.
        typography_applies = zone in ("title", "summary") and id(sh._element) not in grouped
        for para in sh.text_frame.paragraphs:
            ptext = "".join(r.text for r in para.runs)
            if not ptext.strip():
                continue
            sizes = [r.font.size.pt for r in para.runs if r.font.size]
            is_big = zone != "title" and sizes and min(sizes) >= MIN_TITLE_PT
            if is_big:
                big_text.append(ptext)
            elif zone == "summary":
                summary_paras.append(ptext)
                if re.search(r"\d", ptext):
                    numeric_in_summary = True

            if not typography_applies:
                continue
            for run in para.runs:
                if not run.text.strip():
                    continue
                size = run.font.size.pt if run.font.size else None
                if size is None:
                    warn("3 size", f"inherited font size in the {zone} zone; set it explicitly")
                elif zone == "title" and size < MIN_TITLE_PT:
                    fail("3 size", f"title run at {size:g}pt (min {MIN_TITLE_PT})")
                elif zone == "summary" and size < MIN_BODY_PT:
                    fail("3 size", f"summary run at {size:g}pt (min {MIN_BODY_PT})")

                lat, ea = run_fonts(run)
                if CJK_RE.search(run.text):
                    face = ea or lat
                    if face is None:
                        warn("font", "Japanese run with no explicit typeface; set "
                                     f"{sorted(JA_FONTS)[0]}")
                    elif face not in JA_FONTS:
                        fail("font", f"Japanese text rendered in '{face}'; use "
                                     f"{sorted(JA_FONTS)[0]}")
                if LATIN_RE.search(run.text):
                    in_ja_sentence = bool(CJK_RE.search(ptext))
                    if lat is None:
                        warn("font", "Latin run with no explicit typeface; set "
                                     f"{sorted(LATIN_FONTS)[0]}")
                    elif lat not in LATIN_FONTS and in_ja_sentence:
                        warn("font", f"Latin token rendered in '{lat}' inside a Japanese "
                                     f"line; acceptable, but {sorted(LATIN_FONTS)[0]} is "
                                     "the deck's Latin face")
                    elif lat not in LATIN_FONTS:
                        fail("font", f"Latin text rendered in '{lat}'; use "
                                     f"{sorted(LATIN_FONTS)[0]}")

    # --- slide-level composition -------------------------------------------
    # A cover, chapter separator (rule 5) or closing slide: large text, nothing
    # else. These carry no page number and are exempt from the content trio.
    is_cover = bool(big_text) and not summary_paras and not visuals
    if is_cover:
        if pn is not None:
            fail("page number", "cover, chapter separator and closing slides carry no "
                                "page number — remove it")
        return issues

    pn = has_page_number(slide)
    if pn is None:
        fail("page number", "no slide-number field on this slide (a placeholder on the "
                            "layout alone does not print one)")
    else:
        x, y = inches(pn.left), inches(pn.top)
        if x is not None and x < slide_w * 0.6:
            fail("page number", f"not at right (x={x:.2f} in)")
        if y is not None and y > slide_h * 0.15:
            fail("page number", f"not at top (y={y:.2f} in)")
        size, rgb = page_number_style(slide, pn)
        if size is None:
            warn("page number", "page number has no explicit size; set it to "
                                f"{PAGE_NUM_PT}pt on the slide or the master")
        elif size < PAGE_NUM_PT:
            fail("page number", f"page number at {size:g}pt (min {PAGE_NUM_PT})")
        if rgb is not None and rgb.upper() != "000000":
            fail("page number", f"page number in #{rgb}; use pure black")


    if title_sh is None:
        fail("1 composition", "no title at the top of the slide")
    else:
        t = title_sh.text_frame.text.strip().replace("\n", " ")
        if GENERIC_TITLE_RE.match(t):
            fail("4 title", f"generic title '{t}' — use keywords specific to this slide")
        if re.match(r"^\s*\d+(\.\d+)*[\s.:]", t):
            warn("4/5 title", f"title starts with a number ('{t[:24]}…'); a chapter "
                              "separator slide carries that instead")
        if len(t.strip()) < 4:
            warn("4 title", f"title is just '{t}' — looks like an unfinished placeholder")
        if TITLE_TERMINAL_RE.search(t):
            fail("4 title", f"title '{t}' ends like a sentence — use a compact "
                            "keyword phrase; the claim goes in the summary")
        elif TITLE_JA_PREDICATE_RE.search(t):
            fail("4 title", f"title '{t}' ends with a predicate — use a noun phrase "
                            "(e.g. '…の精度'); the claim goes in the summary")
        else:
            m = TITLE_EN_VERB_RE.search(t)
            if m:
                warn("4 title", f"title '{t}' contains '{m.group(0)}' and may be a "
                                "sentence — use a compact keyword phrase")
        n_words = len(re.findall(r"[A-Za-z0-9][\w'’\-]*", CJK_RE.sub(" ", t)))
        n_ja = len(CJK_RE.findall(t))
        if n_words > MAX_TITLE_WORDS:
            warn("4 title", f"title has {n_words} words; compact it to "
                            f"{MAX_TITLE_WORDS} or fewer keywords")
        elif n_ja > MAX_TITLE_JA_CHARS:
            warn("4 title", f"title has {n_ja} Japanese characters; compact it to "
                            f"{MAX_TITLE_JA_CHARS} or fewer")

    if not summary_paras:
        fail("1 composition", "no summary sentences below the title")
    elif len(summary_paras) > 4:
        fail("6 summary", f"{len(summary_paras)} summary lines (max 4; sub-bullets count)")
    for p in summary_paras:
        if len(p.strip()) < 4:
            warn("6 summary", f"summary line is just '{p.strip()}' — unfinished placeholder")
        if len(p) > 220:
            warn("6 summary", f"summary line is {len(p)} chars; split or shorten")
        if re.search(r"\((?:i\.e\.|e\.g\.|note|which means|なお|ただし)", p, re.I) or p.count("(") + p.count("（") >= 3:
            warn("9 supplementary", "parenthetical explanation inside the summary; "
                                    "move it to a caption or balloon")
    if not visuals:
        warn("1 composition", "no figure, table or chart on this slide")
    if title_sh is not None and RESULT_TITLE_RE.search(title_sh.text_frame.text) \
            and not numeric_in_summary:
        warn("7 numbers", "results/data slide with no numeric value in the summary")

    return issues


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--font-en", default="Calibri", help="expected Latin typeface")
    ap.add_argument("--font-ja", default="Yu Gothic", help="expected Japanese typeface")
    args = ap.parse_args()

    global LATIN_FONTS, JA_FONTS
    LATIN_FONTS = {args.font_en}
    JA_FONTS = {args.font_ja} | {args.font_ja + v for v in JA_VARIANTS}

    prs = Presentation(args.pptx)
    slide_w, slide_h = inches(prs.slide_width), inches(prs.slide_height)

    report, n_fail, n_warn = [], 0, 0
    if abs(slide_w - 13.333) > 0.05 or abs(slide_h - 7.5) > 0.05:
        report.append({"slide": 0, "level": "WARN", "rule": "layout",
                       "message": f"canvas is {slide_w:.2f}x{slide_h:.2f} in; the kit assumes 13.33x7.5"})
        n_warn += 1

    for i, slide in enumerate(prs.slides, start=1):
        counts, order = {}, []
        for item in check_slide(slide, slide_w, slide_h):
            if item not in counts:
                order.append(item)
            counts[item] = counts.get(item, 0) + 1
        for item in order:
            level, rule, msg = item
            if counts[item] > 1:
                msg = f"{msg}  [x{counts[item]}]"
            report.append({"slide": i, "level": level, "rule": rule, "message": msg})
            if level == "FAIL":
                n_fail += 1
            else:
                n_warn += 1

    if args.json:
        print(json.dumps({"fails": n_fail, "warns": n_warn, "issues": report},
                         indent=2, ensure_ascii=False))
    else:
        cur = None
        for item in report:
            if item["slide"] != cur:
                cur = item["slide"]
                print(f"\n--- slide {cur} ---")
            print(f"  {item['level']:4} [{item['rule']}] {item['message']}")
        print(f"\n{len(prs.slides)} slide(s) | {n_fail} FAIL, {n_warn} WARN")
        if n_fail == 0 and n_warn == 0:
            print("Mechanical rules pass. Content rules (1, 4, 6, 7, 9) still need a human read.")
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
